from __future__ import annotations

from datetime import datetime
from pathlib import Path
from uuid import uuid4
import base64
import io
import json
import math
import mimetypes
import re

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from flask import Flask, Response, jsonify, render_template, request, send_file, send_from_directory, session
from werkzeug.exceptions import HTTPException
from werkzeug.utils import secure_filename

from config.settings import ARCHIVE_DIR, AUTH_DEFAULT_PASSWORD, AUTH_DEFAULT_PERNER, SECRET_KEY, TRANSFORMED_DIR, UPLOAD_DIR
from services.calculations import MONTH_LABELS, dashboard_summary
from services.database import DatabaseConnectionError
from services.excel_reader import list_sheets, read_excel_file
from services.repository import (
    archive_stored_names,
    authenticate_app_user,
    change_app_user_password,
    create_report_progress,
    delete_employee_master_by_perners,
    delete_report_progress,
    delete_archive_document,
    delete_upload_batch,
    ensure_app_user,
    get_archive_document,
    get_archive_documents,
    get_archive_filter_options,
    get_dashboard_data,
    get_employee_master,
    get_employee_master_meta,
    get_records_by_batch,
    get_report_progress,
    get_upload_batches,
    insert_valid_records,
    lookup_employee,
    normalize_employee_master_dataframe,
    replace_batch_records,
    report_progress_stats,
    save_archive_document,
    save_employee_master_batch,
    save_upload_batch,
    save_validation_errors,
    sync_employee_master_changes,
    upsert_employee_master,
    update_report_progress_status,
    update_report_progress_note,
    validate_employee_master_dataframe,
)
from services.transformer import transform_sanction_data
from services.validator import validate_structure, validate_transformed_data

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 200 * 1024 * 1024
app.config["SECRET_KEY"] = SECRET_KEY
app.config["SESSION_COOKIE_HTTPONLY"] = True
app.config["SESSION_COOKIE_SAMESITE"] = "Lax"


def _is_authenticated() -> bool:
    return bool(session.get("authenticated"))


@app.before_request
def require_login():
    allowed_paths = {"/", "/api/auth/login", "/api/auth/logout", "/api/auth/me"}
    if request.path in allowed_paths or request.path.startswith(("/static/", "/assets/")):
        return None
    if request.path.startswith("/api/") and not _is_authenticated():
        return jsonify({"ok": False, "code": "UNAUTHENTICATED", "message": "Silakan login terlebih dahulu."}), 401
    return None


@app.after_request
def disable_static_cache(response):
    if request.path.startswith("/static/"):
        response.headers["Cache-Control"] = "no-store, no-cache, must-revalidate, max-age=0"
        response.headers["Pragma"] = "no-cache"
        response.headers["Expires"] = "0"
    return response

STAGING_DIR = UPLOAD_DIR / "staging"
MASTER_STAGING_DIR = UPLOAD_DIR / "master_staging"
ARCHIVE_META = ARCHIVE_DIR / "metadata.json"
LATEST_META = TRANSFORMED_DIR / "latest_upload_meta.json"
for directory in (UPLOAD_DIR, TRANSFORMED_DIR, STAGING_DIR, MASTER_STAGING_DIR, ARCHIVE_DIR):
    directory.mkdir(parents=True, exist_ok=True)

SANCTION_DETAIL_SHEET_PATTERN = re.compile(r"^\s*rekap\s+sanksi\s+jan\s*[-–—]\s*des\s+(?:19|20)\d{2}\s*$", re.IGNORECASE)


def _is_sanction_detail_sheet(sheet_name: object) -> bool:
    return bool(SANCTION_DETAIL_SHEET_PATTERN.match(str(sheet_name or "")))

@app.errorhandler(DatabaseConnectionError)
def handle_database_error(exc: DatabaseConnectionError):
    app.logger.warning("Database belum siap: %s", exc)
    return jsonify({
        "ok": False,
        "code": "DATABASE_NOT_READY",
        "message": str(exc),
        "detail": "Pastikan MySQL XAMPP berjalan di host/port sesuai .env, lalu refresh halaman.",
    }), 503


@app.errorhandler(Exception)
def handle_unexpected_error(exc: Exception):
    if isinstance(exc, HTTPException):
        return exc
    app.logger.exception("Unhandled application error")
    return jsonify({"ok": False, "message": str(exc) or "Terjadi kesalahan pada server."}), 500

@app.get("/assets/<path:filename>")
def assets_file(filename: str):
    return send_from_directory("assets", filename)
PDF_STATUS_ORDER = ["Aktif", "Berakhir"]
PDF_STATUS_COLORS = {"Aktif": "#EF4444", "Berakhir": "#007A3D"}
PDF_DEFAULT_CHART_COLORS = ["#007A3D", "#009A4D", "#12A85C", "#35B46D", "#EF4444", "#F87171", "#FCA5A5", "#CBD5E1"]

ALLOWED_ARCHIVE_EXTENSIONS = {
    "pdf", "xlsx", "xls", "xlsm", "xlsb", "xltx", "xltm", "csv", "png", "jpg", "jpeg", "gif", "webp", "doc", "docx",
    "ppt", "pptx", "txt", "md", "json", "xml", "html", "zip", "rar", "7z", "mp4", "mov", "mp3"
}
IMAGE_EXTENSIONS = {"png", "jpg", "jpeg", "gif", "webp"}
EXCEL_EXTENSIONS = {"xlsx", "xls", "xlsm", "xlsb", "xltx", "xltm"}
TABLE_EXTENSIONS = {*EXCEL_EXTENSIONS, "csv"}
TEXT_EXTENSIONS = {"txt", "md", "json", "xml", "html"}


def _json_default(value):
    if pd.isna(value):
        return None
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return str(value)


def _format_duration_days(value: object) -> str:
    try:
        numeric = float(value or 0)
    except (TypeError, ValueError):
        numeric = 0.0
    formatted = f"{numeric:,.1f}".replace(",", "_").replace(".", ",").replace("_", ".")
    return f"{formatted} hari/surat"


def _records(df: pd.DataFrame, limit: int | None = None) -> list[dict[str, object]]:
    if df is None or df.empty:
        return []

    data = df.head(limit).copy() if limit else df.copy()
    for column in ["tanggal_mulai_sanksi", "tanggal_berakhir_sanksi", "tanggal_laporan"]:
        if column in data.columns:
            parsed = pd.to_datetime(data[column], errors="coerce")
            data[column] = [
                parsed_value.strftime("%Y-%m-%d") if not pd.isna(parsed_value) else ("" if pd.isna(original) else str(original))
                for original, parsed_value in zip(data[column], parsed)
            ]

    return json.loads(
        data.to_json(
            orient="records",
            date_format="iso"
        )
    )


def _excel_sheet_name(name: str, used: set[str]) -> str:
    invalid_chars = set('[]:*?/\\')
    clean = "".join("_" if char in invalid_chars else char for char in str(name or "Sheet")).strip() or "Sheet"
    clean = clean[:31]
    base = clean
    counter = 2
    while clean.lower() in used:
        suffix = f"_{counter}"
        clean = f"{base[:31 - len(suffix)]}{suffix}"
        counter += 1
    used.add(clean.lower())
    return clean


def _read_table_file(path: Path) -> pd.DataFrame:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return pd.read_csv(path, dtype=str, keep_default_na=False)
    if suffix.lstrip(".") in EXCEL_EXTENSIONS:
        return pd.read_excel(path, dtype=str, keep_default_na=False)
    raise ValueError("Format file master harus Excel atau CSV.")



def _read_master_workbook(path: Path) -> tuple[pd.DataFrame, pd.DataFrame, list[str]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        return pd.read_csv(path, dtype=str, keep_default_na=False), pd.DataFrame(), ["CSV"]
    if suffix.lstrip(".") in EXCEL_EXTENSIONS:
        excel = pd.ExcelFile(path)
        employee_df = pd.read_excel(path, sheet_name=excel.sheet_names[0], dtype=str, keep_default_na=False) if excel.sheet_names else pd.DataFrame()
        return employee_df, pd.DataFrame(), excel.sheet_names
    raise ValueError("Format file master harus Excel atau CSV.")

def _employee_master_payload(df: pd.DataFrame, errors: list[dict[str, object]] | None = None) -> dict[str, object]:
    preview = df.head(50).where(pd.notna(df), "").to_dict(orient="records") if not df.empty else []
    return {
        "columns": ["perner_karyawan", "nama_karyawan", "jabatan_karyawan", "departemen_seksi"],
        "rows": preview,
        "total": int(len(df)),
        "errors": errors or [],
    }

def _read_latest_transformed() -> pd.DataFrame:
    try:
        db_df = get_dashboard_data()
        if not db_df.empty:
            return db_df
        if get_upload_batches(limit=1).empty:
            return pd.DataFrame()
    except Exception:
        pass

    path = TRANSFORMED_DIR / "latest_transformed.xlsx"
    if not path.exists():
        return pd.DataFrame()
    df = pd.read_excel(path)
    validation = validate_transformed_data(df)
    return validation.valid_df.copy()

def _read_latest_all(batch_id: int | None = None) -> pd.DataFrame:
    try:
        if batch_id is not None:
            batch_df = get_records_by_batch(batch_id)
            return batch_df.drop(columns=["upload_batch_id"], errors="ignore") if not batch_df.empty else pd.DataFrame()
        batches = get_upload_batches(limit=1)
        if batches.empty:
            return pd.DataFrame()
        batch_df = get_records_by_batch(int(batches.iloc[0]["id"]))
        if not batch_df.empty:
            return batch_df.drop(columns=["upload_batch_id"], errors="ignore")
        return pd.DataFrame()
    except Exception:
        pass

    path = TRANSFORMED_DIR / "latest_transformed.xlsx"
    if not path.exists():
        return pd.DataFrame()
    return pd.read_excel(path)


def _is_phk_type(value: object) -> bool:
    text = str(value or "").strip().upper()
    return "PHK" in text or "DISKUALIFIKASI" in text or "MENGUNDURKAN" in text


def _apply_filters(df: pd.DataFrame, filters: dict[str, str]) -> pd.DataFrame:
    if df.empty:
        return df
    result = df.copy()
    if filters.get("year") and filters["year"] != "all" and "tahun_sanksi" in result:
        result = result[result["tahun_sanksi"].astype("Int64") == int(filters["year"])]
    if filters.get("month") and filters["month"] != "all" and "bulan_sanksi" in result:
        result = result[result["bulan_sanksi"].astype("Int64") == int(filters["month"])]
    if filters.get("department") and filters["department"] != "all" and "departemen_seksi" in result:
        result = result[result["departemen_seksi"].astype(str) == filters["department"]]
    if filters.get("type") and filters["type"] != "all" and "jenis_sanksi" in result:
        if str(filters["type"]).strip().upper() == "PHK":
            result = result[result["jenis_sanksi"].map(_is_phk_type)]
        else:
            result = result[result["jenis_sanksi"].astype(str) == filters["type"]]
    if filters.get("position") and filters["position"] != "all" and "jabatan_karyawan" in result:
        result = result[result["jabatan_karyawan"].astype(str) == filters["position"]]
    if filters.get("status") and filters["status"] != "all" and "status_sanksi" in result:
        result = result[result["status_sanksi"].astype(str) == filters["status"]]
    if filters.get("search") and str(filters["search"]).strip() and "perner_karyawan" in result:
        keyword = str(filters["search"]).strip().lower()
        result = result[result["perner_karyawan"].fillna("").astype(str).str.lower().str.contains(keyword, regex=False)]
    return result


def _filter_options(df: pd.DataFrame) -> dict[str, object]:
    def unique(column: str) -> list[str]:
        if df.empty or column not in df:
            return []
        return sorted([str(value) for value in df[column].dropna().unique() if str(value).strip()])

    years = []
    if not df.empty and "tahun_sanksi" in df:
        years = sorted([int(value) for value in df["tahun_sanksi"].dropna().unique()], reverse=True)
    types = unique("jenis_sanksi")
    type_options = []
    seen_types = set()
    for value in types:
        option_value = "PHK" if _is_phk_type(value) else value
        if option_value in seen_types:
            continue
        seen_types.add(option_value)
        type_options.append(option_value)

    return {
        "years": years,
        "months": [{"value": number, "label": label} for number, label in MONTH_LABELS.items()],
        "departments": unique("departemen_seksi"),
        "types": type_options,
        "positions": unique("jabatan_karyawan"),
        "statuses": unique("status_sanksi"),
    }



def _scan_upload_sheets(path: Path) -> tuple[list[dict[str, object]], pd.DataFrame]:
    def validation_details(validation, sheet_name: str) -> list[dict[str, object]]:
        if validation is None or validation.errors_df.empty:
            return []
        details = _records(validation.errors_df, 500)
        for item in details:
            item["sheet"] = sheet_name
        return details

    def valid_details(validation, sheet_name: str) -> list[dict[str, object]]:
        if validation is None or validation.valid_df.empty:
            return []
        columns = [
            "nomor_baris_excel",
            "nomor_surat_sanksi",
            "perner_karyawan",
            "nama_karyawan",
            "departemen_seksi",
            "jenis_sanksi",
            "uraian_pelanggaran",
        ]
        data = validation.valid_df.copy()
        for column in columns:
            if column not in data.columns:
                data[column] = ""
        details = _records(data[columns], 500)
        for item in details:
            item["sheet"] = sheet_name
        return details

    if path.suffix.lower() == ".csv":
        raw_df = pd.read_csv(path, dtype=str, keep_default_na=False)
        structure = validate_structure(raw_df)
        transformed = transform_sanction_data(raw_df) if structure.is_valid else pd.DataFrame()
        validation = validate_transformed_data(transformed) if structure.is_valid else None
        valid_rows = valid_details(validation, "CSV") if structure.is_valid else []
        invalid_details = validation_details(validation, "CSV") if structure.is_valid else [
            {
                "sheet": "CSV",
                "nomor_baris_excel": "",
                "nama_kolom": column,
                "nilai_asli": "",
                "jenis_kesalahan": "Kolom wajib tidak ditemukan",
                "saran_perbaikan": f"Tambahkan atau sesuaikan header kolom {column}.",
            }
            for column in structure.missing_columns
        ]
        return [
            {
                "name": "CSV",
                "valid_structure": structure.is_valid,
                "missing_columns": structure.missing_columns,
                "rows": int(len(raw_df)),
                "valid": 0 if validation is None else int(len(validation.valid_df)),
                "invalid": 0 if validation is None else int(len(validation.invalid_df)),
                "valid_details": valid_rows,
                "invalid_details": invalid_details,
                "selected": structure.is_valid,
            }
        ], raw_df

    scanned = []
    preview_df = pd.DataFrame()
    for sheet_name in list_sheets(path):
        if not _is_sanction_detail_sheet(sheet_name):
            scanned.append(
                {
                    "name": sheet_name,
                    "valid_structure": False,
                    "missing_columns": ["Nama sheet tidak sesuai pola Rekap sanksi Jan - Des YYYY"],
                    "rows": 0,
                    "valid": 0,
                    "invalid": 0,
                    "ignored": True,
                    "ignore_reason": "Nama sheet tidak sesuai pola Rekap sanksi Jan - Des YYYY",
                    "valid_details": [],
                    "invalid_details": [],
                    "selected": False,
                }
            )
            continue
        try:
            raw_df = read_excel_file(path, sheet_name)
            structure = validate_structure(raw_df)
            transformed = transform_sanction_data(raw_df) if structure.is_valid else pd.DataFrame()
            validation = validate_transformed_data(transformed) if structure.is_valid else None
            valid_rows = valid_details(validation, sheet_name) if structure.is_valid else []
            invalid_details = validation_details(validation, sheet_name) if structure.is_valid else [
                {
                    "sheet": sheet_name,
                    "nomor_baris_excel": "",
                    "nama_kolom": column,
                    "nilai_asli": "",
                    "jenis_kesalahan": "Kolom wajib tidak ditemukan",
                    "saran_perbaikan": f"Tambahkan atau sesuaikan header kolom {column}.",
                }
                for column in structure.missing_columns
            ]
            if preview_df.empty and structure.is_valid:
                preview_df = raw_df
            scanned.append(
                {
                    "name": sheet_name,
                    "valid_structure": structure.is_valid,
                    "missing_columns": structure.missing_columns,
                    "rows": int(len(raw_df)),
                    "valid": 0 if validation is None else int(len(validation.valid_df)),
                    "invalid": 0 if validation is None else int(len(validation.invalid_df)),
                    "valid_details": valid_rows,
                    "invalid_details": invalid_details,
                    "selected": structure.is_valid,
                }
            )
        except Exception as exc:
            scanned.append(
                {
                    "name": sheet_name,
                    "valid_structure": False,
                    "missing_columns": [str(exc)],
                    "rows": 0,
                    "valid": 0,
                    "invalid": 0,
                    "valid_details": [],
                    "invalid_details": [
                        {
                            "sheet": sheet_name,
                            "nomor_baris_excel": "",
                            "nama_kolom": "",
                            "nilai_asli": "",
                            "jenis_kesalahan": "Sheet gagal dibaca",
                            "saran_perbaikan": str(exc),
                        }
                    ],
                    "selected": False,
                }
            )
    return scanned, preview_df


def _read_selected_sheets(path: Path, selected_sheets: list[str]) -> tuple[pd.DataFrame, str]:
    if path.suffix.lower() == ".csv":
        return pd.read_csv(path, dtype=str, keep_default_na=False), "CSV"
    frames = []
    loaded_names = []
    for sheet_name in selected_sheets:
        if not _is_sanction_detail_sheet(sheet_name):
            continue
        raw_df = read_excel_file(path, sheet_name)
        structure = validate_structure(raw_df)
        if not structure.is_valid:
            continue
        raw_df = raw_df.copy()
        raw_df["sumber_sheet"] = sheet_name
        frames.append(raw_df)
        loaded_names.append(sheet_name)
    if not frames:
        raise ValueError("Tidak ada sheet detail sanksi yang valid untuk diproses.")
    return pd.concat(frames, ignore_index=True), ", ".join(loaded_names)
def _summary_payload(df: pd.DataFrame) -> dict[str, object]:
    summary = dashboard_summary(df)
    jenis_rows = _records(summary["jenis"])
    departemen_rows = _records(summary["departemen"])
    pelanggar_rows = _records(summary["pelanggar_berulang"])
    kategori_rows = _records(summary["kategori_pelanggaran"])

    def top_value(rows: list[dict[str, object]], key: str) -> str:
        if not rows:
            return "-"
        value = rows[0].get(key)
        text = "" if value is None or pd.isna(value) else str(value).strip()
        return text or "-"

    def status_sanction_notes() -> list[dict[str, object]]:
        if df.empty or "status_sanksi" not in df.columns:
            return []
        detail = df[df["status_sanksi"].astype(str).str.strip().str.lower().isin(["aktif", "berakhir"])].copy()
        if detail.empty:
            return []
        detail["_status_order"] = detail["status_sanksi"].astype(str).str.strip().str.lower().map({"aktif": 0, "berakhir": 1}).fillna(9)
        if "tanggal_mulai_sanksi" in detail.columns:
            detail["_sort_date"] = pd.to_datetime(detail["tanggal_mulai_sanksi"], errors="coerce")
            detail = detail.sort_values(["_status_order", "_sort_date"], ascending=[True, False])
        rows = []
        for _, row in detail.head(10).iterrows():
            start = pd.to_datetime(row.get("tanggal_mulai_sanksi"), errors="coerce")
            end = pd.to_datetime(row.get("tanggal_berakhir_sanksi"), errors="coerce")
            duration = row.get("durasi_sanksi_hari")
            if pd.isna(duration) and not pd.isna(start) and not pd.isna(end):
                duration = max((end.date() - start.date()).days + 1, 0)
            duration_label = ""
            if not pd.isna(duration):
                try:
                    duration_label = f"{int(round(float(duration)))} hari"
                except (TypeError, ValueError):
                    duration_label = str(duration)
            status = str(row.get("status_sanksi") or "").strip().lower()
            status_label = "Aktif" if status == "aktif" else "Berakhir" if status == "berakhir" else str(row.get("status_sanksi") or "")
            rows.append({
                "nomor_surat_sanksi": "" if pd.isna(row.get("nomor_surat_sanksi")) else str(row.get("nomor_surat_sanksi") or ""),
                "status_sanksi": status_label,
                "jenis_sanksi": "" if pd.isna(row.get("jenis_sanksi")) else str(row.get("jenis_sanksi") or ""),
                "uraian_pelanggaran": "" if pd.isna(row.get("uraian_pelanggaran")) else str(row.get("uraian_pelanggaran") or ""),
                "tanggal_mulai_sanksi": "" if pd.isna(start) else start.strftime("%d %b %Y"),
                "tanggal_berakhir_sanksi": "" if pd.isna(end) else end.strftime("%d %b %Y"),
                "durasi_sanksi_hari": "" if pd.isna(duration) else duration,
                "durasi_label": duration_label,
            })
        return rows
    status_notes = status_sanction_notes()

    return {
        "kpi": {
            "total_sanksi": summary["total_sanksi"],
            "total_karyawan": summary["total_karyawan"],
            "total_aktif": summary["total_aktif"],
            "total_pelanggar_berulang": summary["total_pelanggar_berulang"],
            "rata_durasi": round(float(summary["rata_durasi"]), 1),
            "top_jenis_sanksi": top_value(jenis_rows, "jenis_sanksi"),
            "top_departemen": top_value(departemen_rows, "departemen_seksi"),
            "top_pelanggar_berulang": top_value(pelanggar_rows, "nama_karyawan"),
            "top_kategori_pelanggaran": top_value(kategori_rows, "uraian_pelanggaran"),
        },
        "charts": {
            "monthly": _records(summary["bulanan"]),
            "yearly": _records(summary["tahunan"]),
            "types": jenis_rows,
            "departments": departemen_rows,
            "repeatOffenders": pelanggar_rows,
            "violations": kategori_rows,
            "statuses": _records(summary["status"]),
            "statusSanctions": status_notes,
            "activeSanctions": [row for row in status_notes if str(row.get("status_sanksi") or "").lower() == "aktif"],
        },
    }


def _persist_latest(raw_df: pd.DataFrame, transformed_df: pd.DataFrame, source_path: Path, file_name: str, sheet_name: str) -> dict[str, int]:
    batch_file_id = f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uuid4().hex[:8]}"
    batch_dir = TRANSFORMED_DIR / "batches"
    batch_dir.mkdir(parents=True, exist_ok=True)
    raw_path = batch_dir / f"{batch_file_id}_raw.xlsx"
    transformed_path = batch_dir / f"{batch_file_id}_transformed.xlsx"
    raw_df.to_excel(raw_path, index=False)
    transformed_df.to_excel(transformed_path, index=False)

    validation = validate_transformed_data(transformed_df)
    batch_id = save_upload_batch(
        file_name=file_name,
        sheet_name=sheet_name,
        total_rows=len(transformed_df),
        valid_rows=len(validation.valid_df),
        invalid_rows=len(validation.invalid_df),
        duplicate_rows=len(validation.duplicate_df),
        mode="append",
    )
    inserted = insert_valid_records(validation.valid_df, upload_batch_id=batch_id, mode="ignore_duplicates")
    save_validation_errors(validation.errors_df, upload_batch_id=batch_id)
    return {"batch_id": batch_id, "inserted": inserted, "valid": len(validation.valid_df), "invalid": len(validation.invalid_df)}

def _migrate_archive_json_to_mysql() -> None:
    if not ARCHIVE_META.exists():
        return
    try:
        existing_names = archive_stored_names()
        records = json.loads(ARCHIVE_META.read_text(encoding="utf-8"))
        if not isinstance(records, list):
            return
        for record in records:
            stored_name = str(record.get("stored_name") or "")
            if not stored_name or stored_name in existing_names:
                continue
            path = ARCHIVE_DIR / stored_name
            if not path.exists():
                continue
            uploaded_at = None
            try:
                uploaded_at = datetime.fromisoformat(str(record.get("uploaded_at")))
            except Exception:
                pass
            save_archive_document(
                title=str(record.get("title") or "Tanpa Judul"),
                file_name=str(record.get("file_name") or stored_name),
                stored_name=stored_name,
                extension=str(record.get("extension") or Path(stored_name).suffix.lower().lstrip(".")),
                mime_type=str(record.get("mime_type") or mimetypes.guess_type(stored_name)[0] or "application/octet-stream"),
                size=int(record.get("size") or path.stat().st_size),
                uploaded_at=uploaded_at,
            )
    except Exception:
        return


def _archive_by_id(archive_id: str) -> dict[str, object] | None:
    _migrate_archive_json_to_mysql()
    return get_archive_document(archive_id)


@app.get("/")
def index():
    return render_template("index.html")


@app.get("/api/auth/me")
def auth_me_api():
    return jsonify(
        {
            "ok": True,
            "authenticated": _is_authenticated(),
            "user": {"perner": session.get("perner")} if _is_authenticated() else None,
        }
    )


@app.post("/api/auth/login")
def auth_login_api():
    payload = request.get_json(silent=True) or {}
    perner = str(payload.get("perner") or "").strip()
    password = str(payload.get("password") or "")
    ensure_app_user(AUTH_DEFAULT_PERNER, AUTH_DEFAULT_PASSWORD)
    user = authenticate_app_user(perner, password)
    if user:
        session.clear()
        session["authenticated"] = True
        session["perner"] = user["perner"]
        return jsonify({"ok": True, "user": {"perner": user["perner"]}})
    return jsonify({"ok": False, "message": "PERNER atau password salah."}), 401


@app.post("/api/auth/change-password")
def auth_change_password_api():
    if not _is_authenticated():
        return jsonify({"ok": False, "code": "UNAUTHENTICATED", "message": "Silakan login terlebih dahulu."}), 401
    payload = request.get_json(silent=True) or {}
    old_password = str(payload.get("old_password") or "")
    new_password = str(payload.get("new_password") or "")
    confirm_password = str(payload.get("confirm_password") or "")
    if not (1 <= len(new_password) <= 8):
        return jsonify({"ok": False, "message": "Password baru wajib 1 sampai 8 karakter."}), 400
    if new_password != confirm_password:
        return jsonify({"ok": False, "message": "Konfirmasi password tidak sama."}), 400
    if not change_app_user_password(str(session.get("perner") or ""), old_password, new_password):
        return jsonify({"ok": False, "message": "Password lama salah."}), 400
    return jsonify({"ok": True})


@app.post("/api/auth/logout")
def auth_logout_api():
    session.clear()
    return jsonify({"ok": True})




def _pdf_safe_text(value: object, limit: int = 42) -> str:
    text = "" if value is None or pd.isna(value) else str(value)
    text = " ".join(text.split())
    return text if len(text) <= limit else f"{text[: max(0, limit - 1)]}..."


def _draw_pdf_fit_text(canvas, value: object, x: float, y: float, max_width: float, font_name: str = "Helvetica-Bold", font_size: float = 11, min_size: float = 6.5) -> None:
    text = _pdf_safe_text(value, 40)
    size = font_size
    while size > min_size and canvas.stringWidth(text, font_name, size) > max_width:
        size -= .5
    while text and canvas.stringWidth(text, font_name, size) > max_width:
        text = f"{text[:-4].rstrip()}..."
    canvas.setFont(font_name, size)
    canvas.drawString(x, y, text or "-")


def _pdf_chart_rows(payload: dict[str, object], chart_key: str) -> list[dict[str, object]]:
    charts = payload.get("charts", {}) if isinstance(payload, dict) else {}
    rows = charts.get(chart_key, []) if isinstance(charts, dict) else []
    return rows if isinstance(rows, list) else []


def _pdf_value(row: dict[str, object], key: str) -> float:
    try:
        value = row.get(key, 0)
        if value is None or pd.isna(value):
            return 0
        return float(value)
    except Exception:
        return 0


def _draw_pdf_panel(canvas, x: float, top: float, width: float, height: float, title: str) -> tuple[float, float, float, float]:
    from reportlab.lib.colors import HexColor

    bottom = top - height
    canvas.setFillColor(HexColor("#FFFFFF"))
    canvas.setStrokeColor(HexColor("#D7DEE8"))
    canvas.setLineWidth(.65)
    canvas.roundRect(x, bottom, width, height, 5, fill=1, stroke=1)
    canvas.setFillColor(HexColor("#007A3D"))
    canvas.setFont("Helvetica-Bold", 7.2)
    canvas.drawString(x + 8, top - 14, title)
    return x + 8, bottom + 8, width - 16, height - 29
def _draw_pdf_kpis(canvas, payload: dict[str, object], x: float, top: float, width: float) -> None:
    from reportlab.lib.colors import HexColor

    kpi = payload.get("kpi", {}) if isinstance(payload, dict) else {}
    items = [
        ("Total Surat Sanksi", kpi.get("total_sanksi", 0)),
        ("Top Jenis Sanksi", kpi.get("top_jenis_sanksi", "-")),
        ("Top Sanksi Berdasarkan Dept", kpi.get("top_departemen", "-")),
        ("Top Pelanggar Berulang", kpi.get("top_pelanggar_berulang", "-")),
        ("Kategori Pelanggaran Terbanyak", kpi.get("top_kategori_pelanggaran", "-")),
    ]
    gap = 8
    card_w = (width - gap * (len(items) - 1)) / len(items)
    for index, (label, value) in enumerate(items):
        left = x + index * (card_w + gap)
        canvas.setFillColor(HexColor("#FFFFFF"))
        canvas.setStrokeColor(HexColor("#D7DEE8"))
        canvas.roundRect(left, top - 48, card_w, 48, 6, fill=1, stroke=1)
        canvas.setFillColor(HexColor("#64748B"))
        canvas.setFont("Helvetica", 8)
        canvas.drawString(left + 10, top - 17, label)
        canvas.setFillColor(HexColor("#007A3D"))
        canvas.setFont("Helvetica-Bold", 15)
        canvas.drawString(left + 10, top - 37, _pdf_safe_text(value, 18))


def _draw_pdf_line_chart(canvas, rows: list[dict[str, object]], label_key: str, value_key: str, x: float, y: float, width: float, height: float) -> None:
    from reportlab.lib.colors import HexColor

    values = [_pdf_value(row, value_key) for row in rows]
    labels = [_pdf_safe_text(row.get(label_key), 10) for row in rows]
    if not values:
        canvas.setFillColor(HexColor("#64748B"))
        canvas.drawString(x, y + height / 2, "Belum ada data")
        return
    max_value = max(max(values), 1)
    plot_left, plot_bottom = x + 20, y + 20
    plot_w, plot_h = width - 28, height - 34
    canvas.setFont("Helvetica", 5.8)
    canvas.setStrokeColor(HexColor("#E5E7EB"))
    canvas.setLineWidth(.35)
    for i in range(5):
        tick = max_value * i / 4
        gy = plot_bottom + (plot_h * i / 4)
        canvas.line(plot_left, gy, plot_left + plot_w, gy)
        canvas.setFillColor(HexColor("#64748B"))
        canvas.drawRightString(plot_left - 4, gy - 2, str(int(round(tick))))
    points = []
    count = max(len(values) - 1, 1)
    for index, value in enumerate(values):
        px = plot_left + (plot_w * index / count)
        py = plot_bottom + (plot_h * value / max_value)
        points.append((px, py, value))
    canvas.setStrokeColor(HexColor("#007A3D"))
    canvas.setLineWidth(1.5)
    for prev, current in zip(points, points[1:]):
        canvas.line(prev[0], prev[1], current[0], current[1])
    canvas.setFont("Helvetica", 5.8)
    for index, (px, py, value) in enumerate(points):
        canvas.setFillColor(HexColor("#FFFFFF"))
        canvas.setStrokeColor(HexColor("#007A3D"))
        canvas.circle(px, py, 1.8, fill=1, stroke=1)
        canvas.setFillColor(HexColor("#007A3D"))
        canvas.drawCentredString(px, min(py + 8, y + height - 4), str(int(value)))
        if index % max(1, math.ceil(len(labels) / 8)) == 0:
            canvas.setFillColor(HexColor("#64748B"))
            canvas.drawCentredString(px, y + 5, labels[index])
def _draw_pdf_bar_chart(canvas, rows: list[dict[str, object]], label_key: str, value_key: str, x: float, y: float, width: float, height: float) -> None:
    from reportlab.lib.colors import HexColor

    rows = rows[:8]
    values = [_pdf_value(row, value_key) for row in rows]
    if not values:
        canvas.setFillColor(HexColor("#78849A"))
        canvas.drawString(x, y + height / 2, "Belum ada data")
        return
    max_value = max(max(values), 1)
    gap = 8
    bar_w = max(10, (width - 32 - gap * (len(rows) - 1)) / max(len(rows), 1))
    base = y + 24
    plot_h = height - 44
    canvas.setFont("Helvetica", 6.5)
    for index, row in enumerate(rows):
        value = values[index]
        left = x + 24 + index * (bar_w + gap)
        bar_h = plot_h * value / max_value
        canvas.setFillColor(HexColor("#F21B18"))
        canvas.rect(left, base, bar_w, bar_h, fill=1, stroke=0)
        canvas.setFillColor(HexColor("#E8EEF8"))
        canvas.drawCentredString(left + bar_w / 2, base + bar_h + 8, str(int(value)))
        canvas.setFillColor(HexColor("#B8C4D8"))
        canvas.drawCentredString(left + bar_w / 2, y + 8, _pdf_safe_text(row.get(label_key), 9))


def _draw_pdf_hbar_chart(canvas, rows: list[dict[str, object]], label_key: str, value_key: str, x: float, y: float, width: float, height: float, bar_color: str = "#007A3D") -> None:
    from reportlab.lib.colors import HexColor

    rows = rows[:10]
    values = [_pdf_value(row, value_key) for row in rows]
    if not values:
        canvas.setFillColor(HexColor("#64748B"))
        canvas.drawString(x, y + height / 2, "Belum ada data")
        return
    max_value = max(max(values), 1)
    row_h = min(14, (height - 10) / max(len(rows), 1))
    label_w = min(112, width * .40)
    canvas.setFont("Helvetica", 5.9)
    for index, row in enumerate(rows):
        top = y + height - 4 - index * row_h
        value = values[index]
        usable_w = max(12, width - label_w - 25)
        bar_w = usable_w * value / max_value
        canvas.setFillColor(HexColor("#334155"))
        canvas.drawRightString(x + label_w, top - 8, _pdf_safe_text(row.get(label_key), 23))
        canvas.setFillColor(HexColor(bar_color))
        canvas.rect(x + label_w + 7, top - 12, bar_w, 9, fill=1, stroke=0)
        canvas.setFillColor(HexColor("#334155"))
        canvas.drawString(x + label_w + 11 + bar_w, top - 10.5, str(int(value)))
def _draw_pdf_donut(canvas, rows: list[dict[str, object]], label_key: str, value_key: str, x: float, y: float, width: float, height: float) -> None:
    from reportlab.lib.colors import HexColor

    if label_key == "status_sanksi":
        values_by_status = {str(row.get(label_key) or "").strip(): _pdf_value(row, value_key) for row in rows}
        all_rows = [{label_key: status, value_key: values_by_status.get(status, 0)} for status in PDF_STATUS_ORDER]
        colors = [PDF_STATUS_COLORS[status] for status in PDF_STATUS_ORDER]
    else:
        all_rows = list(rows)
        colors = [PDF_DEFAULT_CHART_COLORS[index % len(PDF_DEFAULT_CHART_COLORS)] for index, _ in enumerate(all_rows)]

    drawable = [(index, row) for index, row in enumerate(all_rows) if _pdf_value(row, value_key) > 0]
    total = sum(_pdf_value(row, value_key) for _, row in drawable)
    if total <= 0:
        canvas.setFillColor(HexColor("#64748B"))
        canvas.drawString(x, y + height / 2, "Belum ada data")
        return

    size = min(height - 8, width * .50)
    cx, cy = x + size / 2 + 5, y + height / 2
    start = 90
    for index, row in drawable:
        angle = 360 * _pdf_value(row, value_key) / total
        canvas.setFillColor(HexColor(colors[index]))
        canvas.setStrokeColor(HexColor("#FFFFFF"))
        canvas.setLineWidth(.6)
        canvas.wedge(cx - size / 2, cy - size / 2, cx + size / 2, cy + size / 2, start, -angle, fill=1, stroke=1)
        start -= angle
    canvas.setFillColor(HexColor("#FFFFFF"))
    canvas.circle(cx, cy, size * .29, fill=1, stroke=0)

    legend_x = x + size + 15
    canvas.setFont("Helvetica", 5.7)
    for index, row in enumerate(all_rows[:8]):
        ly = y + height - 15 - index * 11
        canvas.setFillColor(HexColor(colors[index]))
        canvas.rect(legend_x, ly - 5, 5, 5, fill=1, stroke=0)
        canvas.setFillColor(HexColor("#334155"))
        label = _pdf_safe_text(row.get(label_key), 20)
        canvas.drawString(legend_x + 8, ly - 4, label)
def _draw_pdf_table(canvas, rows: list[dict[str, object]], label_key: str, value_key: str, x: float, y: float, width: float, height: float) -> None:
    from reportlab.lib.colors import HexColor

    canvas.setFillColor(HexColor("#F8FAFC"))
    canvas.rect(x, y + height - 24, width, 20, fill=1, stroke=0)
    canvas.setFont("Helvetica-Bold", 6.6)
    canvas.setFillColor(HexColor("#007A3D"))
    canvas.drawString(x + 6, y + height - 17, "KATEGORI")
    canvas.drawRightString(x + width - 6, y + height - 17, "JUMLAH")
    canvas.setStrokeColor(HexColor("#E2E8F0"))
    canvas.line(x, y + height - 27, x + width, y + height - 27)
    canvas.setFont("Helvetica", 6.8)
    row_y = y + height - 43
    for row in rows[:7]:
        canvas.setFillColor(HexColor("#334155"))
        canvas.drawString(x + 6, row_y, _pdf_safe_text(row.get(label_key), 28))
        canvas.setFillColor(HexColor("#007A3D"))
        canvas.setFont("Helvetica-Bold", 7)
        canvas.drawRightString(x + width - 6, row_y, str(int(_pdf_value(row, value_key))))
        canvas.setFont("Helvetica", 6.8)
        canvas.setStrokeColor(HexColor("#E2E8F0"))
        canvas.line(x, row_y - 8, x + width, row_y - 8)
        row_y -= 17


def _draw_pdf_active_sanctions(canvas, rows: list[dict[str, object]], x: float, y: float, width: float, height: float) -> None:
    from reportlab.lib.colors import HexColor

    canvas.setFillColor(HexColor("#F8FAFC"))
    canvas.rect(x, y + height - 24, width, 19, fill=1, stroke=0)
    headers = [
        ("STATUS / JENIS", x + 6, width * .20),
        ("PERIODE", x + width * .24, width * .35),
        ("DURASI", x + width * .61, width * .12),
        ("KETERANGAN", x + width * .75, width * .23),
    ]
    canvas.setFillColor(HexColor("#007A3D"))
    canvas.setFont("Helvetica-Bold", 6.2)
    for label, left, _ in headers:
        canvas.drawString(left, y + height - 17, label)

    if not rows:
        canvas.setFillColor(HexColor("#64748B"))
        canvas.setFont("Helvetica", 7)
        canvas.drawString(x + 6, y + height - 42, "Tidak ada detail sanksi aktif atau berakhir pada filter ini.")
        return

    row_y = y + height - 39
    row_h = 20
    canvas.setFont("Helvetica", 5.8)
    for row in rows[:4]:
        start = _pdf_safe_text(row.get("tanggal_mulai_sanksi"), 12)
        end = _pdf_safe_text(row.get("tanggal_berakhir_sanksi"), 12)
        duration = _pdf_safe_text(row.get("durasi_label") or row.get("durasi_sanksi_hari"), 12)
        if duration and duration.replace(".", "", 1).isdigit():
            duration = f"{duration} hari"
        period = f"{start} - {end}" if end else f"{start} - Aktif"
        status_type = f"{_pdf_safe_text(row.get('status_sanksi'), 12)} / {_pdf_safe_text(row.get('jenis_sanksi'), 14)}"
        values = [
            (status_type, x + 6, width * .20),
            (_pdf_safe_text(period, 40), x + width * .24, width * .35),
            (_pdf_safe_text(duration or "-", 16), x + width * .61, width * .12),
            (_pdf_safe_text(row.get("uraian_pelanggaran"), 42), x + width * .75, width * .23),
        ]
        canvas.setStrokeColor(HexColor("#E2E8F0"))
        canvas.line(x, row_y - 7, x + width, row_y - 7)
        for value, left, max_w in values:
            canvas.setFillColor(HexColor("#334155"))
            _draw_pdf_fit_text(canvas, value, left, row_y, max_w - 4, font_name="Helvetica", font_size=5.8, min_size=4.8)
        row_y -= row_h
def _dashboard_pdf_response(payload: dict[str, object], filters: dict[str, str], row_count: int, as_attachment: bool = True) -> Response:
    try:
        from reportlab.lib.colors import HexColor
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.utils import ImageReader
        from reportlab.pdfgen import canvas
    except ImportError:
        return jsonify({"ok": False, "message": "Library reportlab belum terpasang. Jalankan: pip install reportlab"}), 500

    generated_at = datetime.now()
    buffer = io.BytesIO()
    page_w, page_h = A4
    pdf = canvas.Canvas(buffer, pagesize=A4)

    margin = 28
    content_w = page_w - margin * 2
    green = HexColor("#007A3D")
    red = HexColor("#EF4444")
    muted = HexColor("#64748B")

    year_filter = str(filters.get("year") or filters.get("tahun") or "all")
    period_label = f"Tahun {year_filter}" if year_filter != "all" else "Tahun Semua"
    filter_parts = []
    filter_labels = {
        "year": "Tahun",
        "month": "Bulan",
        "department": "Dept.",
        "type": "Jenis",
        "position": "Jabatan",
        "status": "Status",
        "search": "PERNER",
    }
    for key, label in filter_labels.items():
        value = filters.get(key)
        if value and value != "all":
            filter_parts.append(f"{label}: {value}")
    filter_text = " | ".join(filter_parts) or "Semua data"

    pdf.setFillColor(HexColor("#EEF4EE"))
    pdf.rect(0, 0, page_w, page_h, fill=1, stroke=0)
    pdf.setFillColor(HexColor("#FFFFFF"))
    pdf.rect(0, 0, page_w, page_h, fill=1, stroke=0)

    logo_path = Path("assets") / "app.png"
    if logo_path.exists():
        pdf.drawImage(ImageReader(str(logo_path)), margin, page_h - 57, 34, 34, preserveAspectRatio=True, mask="auto")
    else:
        pdf.setFillColor(red)
        pdf.roundRect(margin, page_h - 57, 34, 34, 6, fill=1, stroke=0)
    pdf.setFillColor(HexColor("#000000"))
    pdf.setFont("Helvetica-Bold", 29)
    pdf.drawString(margin + 39, page_h - 49, "APP")

    pdf.setStrokeColor(green)
    pdf.setLineWidth(1.0)
    pdf.line(margin, page_h - 72, page_w - margin, page_h - 72)

    pill_w = 92
    pdf.setFillColor(HexColor("#E8F7EF"))
    pdf.setStrokeColor(green)
    pdf.roundRect(page_w - margin - pill_w, page_h - 28, pill_w, 14, 2.5, fill=1, stroke=1)
    pdf.setFillColor(green)
    pdf.setFont("Helvetica-Bold", 8)
    pdf.drawCentredString(page_w - margin - pill_w / 2, page_h - 24, f"Periode: {period_label}")
    pdf.setFillColor(green)
    pdf.setFont("Helvetica", 7)
    pdf.drawRightString(page_w - margin, page_h - 42, f"Generated: {generated_at.strftime('%d %b %Y, %H:%M')}")

    pdf.setFillColor(HexColor("#111111"))
    pdf.setFont("Helvetica-Bold", 9.5)
    pdf.drawString(margin, page_h - 91, "Laporan Visualisasi Data Sanksi")
    pdf.setFillColor(green)
    pdf.setFont("Helvetica-Bold", 9)
    pdf.drawString(margin, page_h - 108, f"Data Sanksi Komprehensif ({period_label})")
    pdf.setFillColor(muted)
    pdf.setFont("Helvetica", 7)
    pdf.drawRightString(page_w - margin, page_h - 108, f"Filter: {filter_text} | Rows: {row_count}")

    kpi = payload.get("kpi", {}) if isinstance(payload, dict) else {}
    kpi_items = [
        ("Total Surat Sanksi", kpi.get("total_sanksi", 0)),
        ("Top Jenis Sanksi", kpi.get("top_jenis_sanksi", "-")),
        ("Top Sanksi Berdasarkan Dept", kpi.get("top_departemen", "-")),
        ("Top Pelanggar Berulang", kpi.get("top_pelanggar_berulang", "-")),
        ("Kategori Pelanggaran Terbanyak", kpi.get("top_kategori_pelanggaran", "-")),
    ]
    gap = 4
    kpi_top = page_h - 123
    kpi_h = 39
    kpi_w = (content_w - gap * 4) / 5
    for index, (label, value) in enumerate(kpi_items):
        left = margin + index * (kpi_w + gap)
        pdf.setFillColor(HexColor("#FFFFFF"))
        pdf.setStrokeColor(HexColor("#D7DEE8"))
        pdf.roundRect(left, kpi_top - kpi_h, kpi_w, kpi_h, 4, fill=1, stroke=1)
        pdf.setFillColor(muted)
        _draw_pdf_fit_text(pdf, label, left + 5, kpi_top - 12, kpi_w - 10, font_name="Helvetica", font_size=5.8, min_size=4.6)
        pdf.setFillColor(green if index in (0, 1, 4) else red)
        value_font_size = 11.2 if index == 0 else 8.4
        _draw_pdf_fit_text(pdf, value, left + 5, kpi_top - 29, kpi_w - 10, font_size=value_font_size, min_size=5.4)

    left_w = content_w * .66
    right_w = content_w - left_w - 6
    y_top = page_h - 169
    area = _draw_pdf_panel(pdf, margin, y_top, left_w, 165, "Tren Sanksi Bulanan")
    _draw_pdf_line_chart(pdf, _pdf_chart_rows(payload, "monthly"), "bulan", "jumlah", *area)
    area = _draw_pdf_panel(pdf, margin + left_w + 6, y_top, right_w, 165, "Komposisi Jenis Sanksi")
    _draw_pdf_donut(pdf, _pdf_chart_rows(payload, "types"), "jenis_sanksi", "jumlah", *area)

    y_top = page_h - 338
    half_w = (content_w - 6) / 2
    area = _draw_pdf_panel(pdf, margin, y_top, half_w, 165, "Sanksi Berdasarkan Departemen")
    _draw_pdf_hbar_chart(pdf, _pdf_chart_rows(payload, "departments"), "departemen_seksi", "jumlah", *area, bar_color="#007A3D")
    area = _draw_pdf_panel(pdf, margin + half_w + 6, y_top, half_w, 165, "Top Pelanggar Berulang")
    _draw_pdf_hbar_chart(pdf, _pdf_chart_rows(payload, "repeatOffenders"), "nama_karyawan", "jumlah", *area, bar_color="#EF4444")

    y_top = page_h - 511
    third_w = (content_w - 12) / 3
    area = _draw_pdf_panel(pdf, margin, y_top, third_w, 170, "Tren Sanksi Tahunan")
    _draw_pdf_line_chart(pdf, _pdf_chart_rows(payload, "yearly"), "tahun_sanksi", "jumlah", *area)
    area = _draw_pdf_panel(pdf, margin + third_w + 6, y_top, third_w, 170, "Status Sanksi")
    _draw_pdf_donut(pdf, _pdf_chart_rows(payload, "statuses"), "status_sanksi", "jumlah", *area)
    area = _draw_pdf_panel(pdf, margin + (third_w + 6) * 2, y_top, third_w, 170, "Kategori Pelanggaran")
    _draw_pdf_table(pdf, _pdf_chart_rows(payload, "violations"), "uraian_pelanggaran", "jumlah", *area)

    has_detail_filter = any(str(filters.get(key) or "").strip() and str(filters.get(key) or "").strip() != "all" for key in filter_labels)
    if has_detail_filter:
        area = _draw_pdf_panel(pdf, margin, 153, content_w, 126, "Keterangan Status Sanksi Berdasarkan Filter")
        _draw_pdf_active_sanctions(pdf, _pdf_chart_rows(payload, "statusSanctions"), *area)

    pdf.setFillColor(green)
    pdf.rect(0, 0, page_w, 13, fill=1, stroke=0)
    pdf.setFillColor(HexColor("#FFFFFF"))
    pdf.setFont("Helvetica-Bold", 6.8)
    pdf.drawCentredString(page_w / 2, 4, f"\u00a9 {generated_at.year} APP Region Jambi")

    pdf.save()
    buffer.seek(0)
    filename = f"dashboard_visual_sanksi_{generated_at.strftime('%Y%m%d_%H%M%S')}.pdf"
    return Response(
        buffer.getvalue(),
        mimetype="application/pdf",
        headers={"Content-Disposition": f"{'attachment' if as_attachment else 'inline'}; filename={filename}"},
    )

@app.get("/api/master-employees")
def master_employees_api():
    search = request.args.get("search") or None
    status = request.args.get("status") or None
    df = get_employee_master(search=search, status=status)
    meta = get_employee_master_meta()
    rows = df.where(pd.notna(df), "").to_dict(orient="records") if not df.empty else []
    return jsonify({"ok": True, "columns": ["perner_karyawan", "nama_karyawan", "jabatan_karyawan", "departemen_seksi"], "rows": rows, "meta": meta})


@app.post("/api/master-employees/preview")
def master_employees_preview_api():
    uploaded = request.files.get("file")
    if uploaded is None:
        return jsonify({"ok": False, "message": "File master wajib dipilih."}), 400
    filename = secure_filename(uploaded.filename or f"master_{uuid4().hex}.xlsx")
    if Path(filename).suffix.lower().lstrip(".") not in TABLE_EXTENSIONS:
        return jsonify({"ok": False, "message": "Format master harus Excel atau CSV."}), 400
    path = MASTER_STAGING_DIR / f"{uuid4().hex}_{filename}"
    uploaded.save(path)
    employee_raw, _, sheet_names = _read_master_workbook(path)
    employee_df, employee_errors = validate_employee_master_dataframe(employee_raw)
    payload = _employee_master_payload(employee_df, employee_errors)
    payload.update({
        "ok": True,
        "path": str(path),
        "file_name": filename,
        "sheet_names": sheet_names,
        "has_errors": bool(employee_errors),
    })
    return jsonify(payload)


@app.post("/api/master-employees/load")
def master_employees_load_api():
    payload = request.get_json(silent=True) or {}
    path = Path(str(payload.get("path", "")))
    mode = str(payload.get("mode") or "update").lower()
    if mode not in {"update", "replace"}:
        mode = "update"
    if not path.exists() or MASTER_STAGING_DIR not in path.parents:
        return jsonify({"ok": False, "message": "File master tidak ditemukan."}), 400
    employee_raw, _, sheet_names = _read_master_workbook(path)
    normalized, errors = validate_employee_master_dataframe(employee_raw)
    if errors:
        return jsonify({"ok": False, "message": "Master belum valid.", "errors": errors}), 400
    batch_id = save_employee_master_batch(str(payload.get("file_name") or path.name), len(normalized), len(normalized), 0, 0, mode)
    saved = upsert_employee_master(normalized, batch_id=batch_id, mode=mode)
    return jsonify({"ok": True, "saved": saved, "sheets": sheet_names, "meta": get_employee_master_meta()})


@app.post("/api/master-employees")
def master_employees_save_api():
    payload = request.get_json(silent=True) or {}
    rows = payload.get("rows", [])
    mode = str(payload.get("mode") or "update").lower()
    df = pd.DataFrame(rows)
    if rows:
        _, errors = validate_employee_master_dataframe(df)
        if errors:
            detail = "; ".join(str(error.get("message") or "") for error in errors[:5] if error.get("message"))
            message = f"Master belum valid. {detail}" if detail else "Master belum valid."
            return jsonify({"ok": False, "message": message, "errors": errors}), 400
    try:
        result = sync_employee_master_changes(df, payload.get("deleted_perners", []), mode)
    except ValueError as exc:
        return jsonify({"ok": False, "message": str(exc)}), 400
    return jsonify({"ok": True, **result, "meta": get_employee_master_meta()})


@app.get("/api/master-employees/lookup")
def master_employees_lookup_api():
    return jsonify({"ok": True, "records": lookup_employee(perner=request.args.get("perner"), name=request.args.get("name"), limit=10)})


@app.get("/api/master-employees/export")
def master_employees_export_api():
    df = get_employee_master()
    employee_columns = ["perner_karyawan", "nama_karyawan", "jabatan_karyawan", "departemen_seksi"]
    employee_export = df.reindex(columns=employee_columns)
    output = io.BytesIO()
    employee_export.to_excel(output, sheet_name="Master Karyawan", index=False)
    output.seek(0)
    return send_file(output, as_attachment=True, download_name=f"master_data_karyawan_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx", mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


@app.get("/api/reports")
def report_progress_list_api():
    search = request.args.get("search")
    status = request.args.get("status")
    filtered = get_report_progress(search=search, status=status)
    return jsonify({"ok": True, "records": filtered, "stats": report_progress_stats(filtered)})


@app.get("/api/reports/export")
def report_progress_export_api():
    search = request.args.get("search")
    status = request.args.get("status")
    records = get_report_progress(search=search, status=status)
    status_labels = {
        "laporan_masuk": "Laporan Masuk",
        "diterima_process": "Diterima / Process",
        "pemanggilan_investigasi": "Pemanggilan & Investigasi",
        "pemeriksaan": "Pemeriksaan",
        "sk_selesai": "Rekom Diterbitkan",
    }
    status_keys = list(status_labels.keys())
    headers = ["NO", "NOMOR LAPORAN", "DEPARTEMEN", "TANGGAL", "JENIS LAPORAN"]
    headers += [f"STATUS: {label}" for label in status_labels.values()]
    headers.append("KETERANGAN")
    wb = Workbook()
    ws = wb.active
    ws.title = "Progres Laporan"
    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="007A3D", end_color="007A3D", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)
    check_font = Font(bold=True, color="007A3D", size=12)
    thin_border = Border(
        left=Side(style="thin", color="D1D5DB"),
        right=Side(style="thin", color="D1D5DB"),
        top=Side(style="thin", color="D1D5DB"),
        bottom=Side(style="thin", color="D1D5DB"),
    )
    for column_index, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=column_index, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        cell.border = thin_border
    ws.freeze_panes = "A2"
    for row_index, record in enumerate(records, start=2):
        tanggal_val = record.get("tanggal_laporan") or ""
        if hasattr(tanggal_val, "strftime"):
            tanggal_display = tanggal_val.strftime("%d/%m/%Y")
        else:
            try:
                tanggal_display = pd.to_datetime(tanggal_val, errors="coerce").strftime("%d/%m/%Y")
            except Exception:
                tanggal_display = str(tanggal_val) if tanggal_val else "-"
        active_status = str(record.get("status_progress") or "laporan_masuk")
        is_final = active_status == "sk_selesai"
        values = [
            row_index - 1,
            record.get("nomor_laporan") or "",
            record.get("departemen") or "",
            tanggal_display,
            record.get("jenis_laporan") or "",
        ]
        for key in status_keys:
            values.append("✓" if is_final or key == active_status else "")
        values.append(record.get("keterangan") or "")
        for column_index, value in enumerate(values, start=1):
            cell = ws.cell(row=row_index, column=column_index, value=value)
            cell.border = thin_border
            cell.alignment = Alignment(vertical="center", wrap_text=(column_index == len(headers)))
            if 6 <= column_index <= 10 and value == "✓":
                cell.font = check_font
                cell.alignment = Alignment(horizontal="center", vertical="center")
            elif column_index == 1:
                cell.alignment = Alignment(horizontal="center", vertical="center")
    column_widths = {
        "A": 6,
        "B": 20,
        "C": 28,
        "D": 16,
        "E": 30,
    }
    status_start_col = 6
    for i, _key in enumerate(status_keys):
        col_letter = ws.cell(row=1, column=status_start_col + i).column_letter
        column_widths[col_letter] = 24
    col_k = ws.cell(row=1, column=len(headers)).column_letter
    column_widths[col_k] = 36
    for col_letter, width in column_widths.items():
        ws.column_dimensions[col_letter].width = width
    ws.auto_filter.ref = ws.dimensions
    output = io.BytesIO()
    wb.save(output)
    output.seek(0)
    filename = f"Progres_Laporan_{datetime.now().strftime('%Y-%m-%d')}.xlsx"
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )


@app.post("/api/reports")
def report_progress_create_api():
    payload = request.get_json(silent=True) or {}
    try:
        record = create_report_progress(payload)
    except ValueError as exc:
        return jsonify({"ok": False, "message": str(exc)}), 400
    all_records = get_report_progress()
    return jsonify({"ok": True, "record": record, "records": all_records, "stats": report_progress_stats(all_records)})


@app.patch("/api/reports/<int:report_id>/status")
def report_progress_status_api(report_id: int):
    payload = request.get_json(silent=True) or {}
    try:
        record = update_report_progress_status(report_id, str(payload.get("status") or ""))
    except ValueError as exc:
        return jsonify({"ok": False, "message": str(exc)}), 400
    all_records = get_report_progress()
    return jsonify({"ok": True, "record": record, "records": all_records, "stats": report_progress_stats(all_records)})


@app.patch("/api/reports/<int:report_id>/note")
def report_progress_note_api(report_id: int):
    payload = request.get_json(silent=True) or {}
    try:
        record = update_report_progress_note(report_id, str(payload.get("keterangan") or ""))
    except ValueError as exc:
        return jsonify({"ok": False, "message": str(exc)}), 400
    all_records = get_report_progress()
    return jsonify({"ok": True, "record": record, "records": all_records, "stats": report_progress_stats(all_records)})


@app.delete("/api/reports/<int:report_id>")
def report_progress_delete_api(report_id: int):
    try:
        delete_report_progress(report_id)
    except ValueError as exc:
        return jsonify({"ok": False, "message": str(exc)}), 400
    all_records = get_report_progress()
    return jsonify({"ok": True, "records": all_records, "stats": report_progress_stats(all_records)})

@app.get("/api/dashboard")
def dashboard_api():
    df = _read_latest_transformed()
    filtered = _apply_filters(df, request.args.to_dict())
    payload = _summary_payload(filtered)
    payload["filters"] = _filter_options(df)
    payload["rows"] = int(len(filtered))
    return jsonify(payload)



@app.get("/api/dashboard/export")
def dashboard_export_api():
    filters = request.args.to_dict()
    as_attachment = filters.pop("download", "1") == "1"
    df = _read_latest_transformed()
    filtered = _apply_filters(df, filters)
    csv_data = filtered.to_csv(index=False, encoding="utf-8-sig")
    filename = f"dashboard_sanksi_{datetime.now().strftime('%Y%m%d_%H%M%S')}.csv"
    return Response(
        csv_data,
        mimetype="text/csv; charset=utf-8",
        headers={"Content-Disposition": f"{'attachment' if as_attachment else 'inline'}; filename={filename}"},
    )


@app.get("/api/dashboard/export/pdf")
def dashboard_pdf_export_api():
    filters = request.args.to_dict()
    as_attachment = filters.pop("download", "0") == "1"
    filters.pop("preview", None)
    df = _read_latest_transformed()
    filtered = _apply_filters(df, filters)
    payload = _summary_payload(filtered)
    return _dashboard_pdf_response(payload, filters, int(len(filtered)), as_attachment=as_attachment)
@app.get("/api/transform")
def transform_api():
    requested_batch = request.args.get("batch_id")
    all_batches = str(requested_batch or "").strip().lower() == "all"
    batch_id = int(requested_batch) if requested_batch and requested_batch not in {"latest", "all"} else None
    df = get_dashboard_data().drop(columns=["upload_batch_id"], errors="ignore") if all_batches else _read_latest_all(batch_id=batch_id)
    validation = validate_transformed_data(df) if not df.empty else None
    batches_df = get_upload_batches(limit=100)
    batches = []
    for row in batches_df.to_dict(orient="records") if not batches_df.empty else []:
        sheet_name = str(row.get("sheet_name") or "Data Sanksi")
        batches.append(
            {
                "id": int(row["id"]),
                "file_name": row.get("file_name"),
                "sheet_name": sheet_name,
                "valid_rows": int(row.get("valid_rows") or 0),
                "created_at": str(row.get("created_at")),
            }
        )
    selected = (
        {
            "id": "all",
            "file_name": "Pilih Semua",
            "sheet_name": "Semua Data Sanksi",
            "valid_rows": int(len(df)),
            "created_at": "",
        }
        if all_batches
        else next((batch for batch in batches if batch_id is not None and batch["id"] == batch_id), batches[0] if batches else {})
    )
    sheets = [sheet.strip() for sheet in str(selected.get("sheet_name") or "Data Sanksi").split(",") if sheet.strip()]
    meta = {**selected, "batch_id": selected.get("id"), "sheets": sheets or ["Data Sanksi"]} if selected else {}
    return jsonify(
        {
            "meta": meta,
            "batches": batches,
            "selected_batch_id": selected.get("id") if selected else None,
            "columns": list(df.columns),
            "rows": _records(df, 1000),
            "valid": 0 if validation is None else len(validation.valid_df),
            "invalid": 0 if validation is None else len(validation.invalid_df),
            "errors": [] if validation is None else _records(validation.errors_df, 200),
        }
    )

@app.delete("/api/transform/<int:batch_id>")
def delete_transform_batch_api(batch_id: int):
    deleted = delete_upload_batch(batch_id)
    if not deleted:
        return jsonify({"ok": False, "message": "File/batch tidak ditemukan."}), 404
    batches_df = get_upload_batches(limit=1)
    next_batch_id = None if batches_df.empty else int(batches_df.iloc[0]["id"])
    valid_df = get_dashboard_data()
    return jsonify(
        {
            "ok": True,
            "deleted": deleted,
            "next_batch_id": next_batch_id,
            "dashboard": _summary_payload(valid_df),
            "filters": _filter_options(valid_df),
        }
    )

@app.delete("/api/transform/all")
def delete_all_transform_batches_api():
    batches_df = get_upload_batches(limit=10000)
    deleted_records = []
    for row in batches_df.to_dict(orient="records") if not batches_df.empty else []:
        deleted = delete_upload_batch(int(row["id"]))
        if deleted:
            deleted_records.append(deleted)
    valid_df = get_dashboard_data()
    return jsonify(
        {
            "ok": True,
            "deleted": deleted_records,
            "deleted_count": len(deleted_records),
            "next_batch_id": None,
            "dashboard": _summary_payload(valid_df),
            "filters": _filter_options(valid_df),
        }
    )
@app.post("/api/transform/export")
def export_transform_excel_api():
    payload = request.get_json(silent=True) or {}
    columns = [str(column) for column in payload.get("columns", []) if str(column).strip()]
    sheets_payload = payload.get("sheets", [])
    if not columns:
        return jsonify({"ok": False, "message": "Kolom tidak tersedia untuk export."}), 400
    if not sheets_payload:
        sheets_payload = [{"name": "Data Sanksi", "rows": payload.get("rows", [])}]

    output = io.BytesIO()
    used_sheet_names: set[str] = set()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        wrote_sheet = False
        for index, sheet in enumerate(sheets_payload, start=1):
            name = _excel_sheet_name(str(sheet.get("name") or f"Sheet {index}"), used_sheet_names)
            rows = sheet.get("rows", []) or []
            df = pd.DataFrame(rows)
            for column in columns:
                if column not in df.columns:
                    df[column] = ""
            df = df[columns].copy()
            df = df.dropna(how="all")
            if not df.empty:
                df = df.loc[~df.apply(lambda row: all(str(value).strip() == "" or str(value).lower() == "nan" for value in row), axis=1)]
            df.to_excel(writer, sheet_name=name, index=False)
            worksheet = writer.sheets[name]
            worksheet.freeze_panes = "A2"
            for cell in worksheet[1]:
                cell.font = cell.font.copy(bold=True)
            for column_cells in worksheet.columns:
                values = [str(cell.value or "") for cell in column_cells]
                width = min(max(12, max(len(value) for value in values) + 2), 42)
                worksheet.column_dimensions[column_cells[0].column_letter].width = width
            wrote_sheet = True
        if not wrote_sheet:
            pd.DataFrame(columns=columns).to_excel(writer, sheet_name="Data Sanksi", index=False)
    output.seek(0)
    filename = f"transform_data_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return send_file(
        output,
        as_attachment=True,
        download_name=filename,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
@app.post("/api/transform")
def save_transform_api():
    payload = request.get_json(silent=True) or {}
    rows = payload.get("rows", [])
    sheet_names = [str(sheet).strip() for sheet in payload.get("sheets", []) if str(sheet).strip()]
    source_file_name = str(payload.get("source_file_name") or "").strip()
    source_sheet_name = str(payload.get("source_sheet_name") or "").strip()
    df = pd.DataFrame(rows)
    if df.empty:
        return jsonify({"ok": False, "message": "Tidak ada data yang bisa disimpan."}), 400

    transformed = transform_sanction_data(df)
    validation = validate_transformed_data(transformed)
    if validation.valid_df.empty:
        return jsonify(
            {
                "ok": False,
                "message": "Tidak ada baris valid untuk disimpan. Periksa kembali kolom wajib dan format tanggal.",
                "errors": _records(validation.errors_df, 500),
                "valid": len(validation.valid_df),
                "invalid": len(validation.invalid_df),
            }
        ), 400

    batch_dir = TRANSFORMED_DIR / "batches"
    batch_dir.mkdir(parents=True, exist_ok=True)
    transformed_file_name = secure_filename(source_file_name or "") or "transform_manual.xlsx"
    transformed.to_excel(batch_dir / f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uuid4().hex[:8]}_{transformed_file_name}", index=False)

    source_batch_id = payload.get("batch_id")
    source_batch = pd.DataFrame()
    if source_batch_id:
        batches = get_upload_batches(limit=100)
        if not batches.empty:
            source_batch = batches[batches["id"].astype(int) == int(source_batch_id)]
    if source_batch.empty and not source_file_name:
        source_batch = get_upload_batches(limit=1)
    latest_meta = {} if source_batch.empty else {"file_name": source_batch.iloc[0]["file_name"], "sheet_name": source_batch.iloc[0]["sheet_name"]}
    if not source_batch.empty and source_batch_id:
        batch_id = int(source_batch.iloc[0]["id"])
        inserted = replace_batch_records(validation.valid_df, upload_batch_id=batch_id)
        save_validation_errors(validation.errors_df, upload_batch_id=batch_id)
    else:
        batch_id = save_upload_batch(
            file_name=source_file_name or str(latest_meta.get("file_name") or "Transform manual"),
            sheet_name=source_sheet_name or (", ".join(sheet_names) if sheet_names else str(latest_meta.get("sheet_name") or "Data Sanksi")),
            total_rows=len(transformed),
            valid_rows=len(validation.valid_df),
            invalid_rows=len(validation.invalid_df),
            duplicate_rows=len(validation.duplicate_df),
            mode="append",
        )
        inserted = insert_valid_records(validation.valid_df, upload_batch_id=batch_id, mode="ignore_duplicates")
        save_validation_errors(validation.errors_df, upload_batch_id=batch_id)
    valid_df = get_dashboard_data()
    return jsonify({"ok": True, "batch_id": batch_id, "inserted": inserted, "valid": len(validation.valid_df), "invalid": len(validation.invalid_df), "dashboard": _summary_payload(valid_df), "filters": _filter_options(valid_df)})

@app.post("/api/sanctions/preview")
def preview_sanctions_api():
    uploaded = request.files.get("file")
    if uploaded is None:
        return jsonify({"ok": False, "message": "File wajib dipilih."}), 400
    filename = secure_filename(uploaded.filename or f"upload_{uuid4().hex}.xlsx")
    upload_id = uuid4().hex
    path = STAGING_DIR / f"{upload_id}_{filename}"
    uploaded.save(path)
    try:
        sheet_details, preview_df = _scan_upload_sheets(path)
        selected_sheets = [sheet["name"] for sheet in sheet_details if sheet["selected"]]
        total_valid = sum(int(sheet["valid"]) for sheet in sheet_details if sheet["selected"])
        total_invalid = sum(int(sheet["invalid"]) for sheet in sheet_details if sheet["selected"])
        return jsonify(
            {
                "ok": True,
                "upload_id": upload_id,
                "path": str(path),
                "file_name": filename,
                "sheet_name": ", ".join(selected_sheets) if selected_sheets else "",
                "selected_sheets": selected_sheets,
                "sheet_details": sheet_details,
                "sheets": [sheet["name"] for sheet in sheet_details],
                "columns": list(preview_df.columns),
                "rows": _records(preview_df, 12),
                "structure_valid": bool(selected_sheets),
                "missing_columns": [] if selected_sheets else ["Tidak ada sheet detail sanksi yang strukturnya valid."],
                "valid": total_valid,
                "invalid": total_invalid,
            }
        )
    except Exception as exc:
        return jsonify({"ok": False, "message": str(exc)}), 400

@app.post("/api/sanctions/load")
def load_sanctions_api():
    payload = request.get_json(silent=True) or {}
    path = Path(str(payload.get("path", "")))
    selected_sheets = payload.get("selected_sheets") or []
    file_name = str(payload.get("file_name") or path.name)
    if not path.exists() or STAGING_DIR not in path.parents:
        return jsonify({"ok": False, "message": "File upload tidak ditemukan."}), 400
    if not selected_sheets:
        sheet_details, _ = _scan_upload_sheets(path)
        selected_sheets = [sheet["name"] for sheet in sheet_details if sheet.get("valid_structure")]
    if not selected_sheets:
        return jsonify({"ok": False, "message": "Pilih minimal satu sheet detail sanksi yang valid. Pastikan header berada pada baris pertama sheet detail sanksi."}), 400
    raw_df, sheet_name = _read_selected_sheets(path, [str(sheet) for sheet in selected_sheets])
    transformed = transform_sanction_data(raw_df)
    validation = validate_transformed_data(transformed)
    final_path = UPLOAD_DIR / f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_{uuid4().hex[:8]}_{file_name}"
    final_path.write_bytes(path.read_bytes())
    persist_result = _persist_latest(raw_df, transformed, final_path, file_name, sheet_name)
    db_df = get_dashboard_data()
    return jsonify({"ok": True, **persist_result, "dashboard": _summary_payload(db_df), "filters": _filter_options(db_df)})

@app.post("/api/sanctions/transform-preview")
def transform_preview_sanctions_api():
    payload = request.get_json(silent=True) or {}
    path = Path(str(payload.get("path") or ""))
    if not path.exists() or STAGING_DIR not in path.parents:
        return jsonify({"ok": False, "message": "File staging tidak ditemukan. Upload ulang file terlebih dahulu."}), 400

    sheet_details, _ = _scan_upload_sheets(path)
    available_sheets = {str(sheet["name"]) for sheet in sheet_details if sheet.get("valid_structure")}
    selected_sheets = [str(sheet) for sheet in payload.get("selected_sheets", []) if str(sheet) in available_sheets]
    if not selected_sheets:
        selected_sheets = [str(sheet["name"]) for sheet in sheet_details if sheet.get("selected")]
    if not selected_sheets:
        return jsonify({"ok": False, "message": "Pilih minimal satu sheet detail sanksi yang valid."}), 400

    raw_df, sheet_name = _read_selected_sheets(path, selected_sheets)
    transformed = transform_sanction_data(raw_df)
    grid_df = transformed.drop(columns=["nomor_baris_excel"], errors="ignore")
    validation = validate_transformed_data(transformed)

    batches_df = get_upload_batches(limit=100)
    batches = []
    for row in batches_df.to_dict(orient="records") if not batches_df.empty else []:
        batches.append(
            {
                "id": int(row["id"]),
                "file_name": row.get("file_name"),
                "sheet_name": row.get("sheet_name"),
                "valid_rows": int(row.get("valid_rows") or 0),
                "created_at": str(row.get("created_at")),
            }
        )

    return jsonify(
        {
            "ok": True,
            "meta": {
                "id": "",
                "batch_id": "",
                "file_name": payload.get("file_name") or path.name,
                "sheet_name": sheet_name,
                "sheets": [sheet.strip() for sheet in str(sheet_name).split(",") if sheet.strip()] or ["Data Sanksi"],
                "valid_rows": int(len(validation.valid_df)),
                "created_at": "",
            },
            "batches": batches,
            "selected_batch_id": "",
            "columns": list(grid_df.columns),
            "rows": _records(grid_df, 5000),
            "valid": len(validation.valid_df),
            "invalid": len(validation.invalid_df),
            "errors": _records(validation.errors_df, 500),
        }
    )

@app.get("/api/archives")
def archives_api():
    _migrate_archive_json_to_mysql()
    year_value = request.args.get("year", "all")
    title_value = request.args.get("title", "all")
    year = None if year_value == "all" else int(year_value)
    title = None if title_value == "all" else title_value
    records = get_archive_documents(year=year, title=title)
    filters = get_archive_filter_options()
    return jsonify({"records": records, "filters": filters})


@app.post("/api/archives")
def archive_upload_api():
    title = (request.form.get("title") or "").strip()
    uploaded = request.files.get("file")
    if not title or uploaded is None:
        return jsonify({"ok": False, "message": "Judul dan file wajib diisi."}), 400
    filename = secure_filename(uploaded.filename or "archive")
    extension = Path(filename).suffix.lower().lstrip(".")
    if extension not in ALLOWED_ARCHIVE_EXTENSIONS:
        return jsonify({"ok": False, "message": "Format file tidak diizinkan."}), 400
    archive_id = uuid4().hex
    stored_name = f"{datetime.now().strftime('%Y%m%d_%H%M%S')}_{archive_id[:8]}.{extension}"
    path = ARCHIVE_DIR / stored_name
    uploaded.save(path)
    archive_db_id = save_archive_document(
        title=title,
        file_name=filename,
        stored_name=stored_name,
        extension=extension,
        mime_type=uploaded.mimetype or mimetypes.guess_type(filename)[0] or "application/octet-stream",
        size=path.stat().st_size,
    )
    return jsonify({"ok": True, "id": archive_db_id})



@app.delete("/api/archives/<archive_id>")
def archive_delete_api(archive_id: str):
    record = _archive_by_id(archive_id)
    if not record:
        return jsonify({"ok": False, "message": "Arsip tidak ditemukan."}), 404
    deleted = delete_archive_document(archive_id)
    if not deleted:
        return jsonify({"ok": False, "message": "Arsip tidak ditemukan."}), 404
    path = ARCHIVE_DIR / str(record.get("stored_name"))
    if path.exists():
        path.unlink()
    return jsonify({"ok": True})

@app.get("/api/archives/<archive_id>/file")
def archive_file_api(archive_id: str):
    record = _archive_by_id(archive_id)
    if not record:
        return jsonify({"message": "Arsip tidak ditemukan."}), 404
    path = ARCHIVE_DIR / str(record.get("stored_name"))
    return send_file(path, mimetype=str(record.get("mime_type") or "application/octet-stream"), as_attachment=False, download_name=str(record.get("file_name") or path.name))


@app.get("/api/archives/<archive_id>/download")
def archive_download_api(archive_id: str):
    record = _archive_by_id(archive_id)
    if not record:
        return jsonify({"message": "Arsip tidak ditemukan."}), 404
    path = ARCHIVE_DIR / str(record.get("stored_name"))
    return send_file(path, mimetype=str(record.get("mime_type") or "application/octet-stream"), as_attachment=True, download_name=str(record.get("file_name") or path.name))


@app.get("/api/archives/<archive_id>/preview")
def archive_preview_api(archive_id: str):
    record = _archive_by_id(archive_id)
    if not record:
        return jsonify({"message": "Arsip tidak ditemukan."}), 404
    path = ARCHIVE_DIR / str(record.get("stored_name"))
    ext = str(record.get("extension", "")).lower()
    if ext in TABLE_EXTENSIONS:
        df = pd.read_csv(path) if ext == "csv" else pd.read_excel(path)
        return jsonify({"kind": "table", "columns": list(df.columns), "rows": _records(df, 300)})
    if ext in TEXT_EXTENSIONS:
        return jsonify({"kind": "text", "text": path.read_text(encoding="utf-8", errors="replace")[:20000]})
    if ext in IMAGE_EXTENSIONS:
        return jsonify({"kind": "image", "url": f"/api/archives/{archive_id}/file"})
    if ext == "pdf":
        return jsonify({"kind": "pdf", "url": f"/api/archives/{archive_id}/file"})
    return jsonify({"kind": "unsupported", "message": "Preview untuk format ini belum didukung. File tetap dapat diunduh."})


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=8501, debug=True)














































